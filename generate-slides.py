#!/usr/bin/env python3
"""Generate 6-page demo slides framework for CP2."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(5.625)

# Color palette: Teal Trust
PRIMARY = RGBColor(2, 128, 144)
SECONDARY = RGBColor(0, 168, 150)
ACCENT = RGBColor(2, 195, 154)
DARK = RGBColor(33, 33, 33)
GRAY = RGBColor(107, 114, 128)
LIGHT_GRAY = RGBColor(243, 244, 246)

def add_title_slide(prs, title_text, content_blocks, bg_color=PRIMARY):
    """Add a title slide with dark background."""
    blank_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(blank_layout)

    # Background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_color

    # Title
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
    tf = title.text_frame
    tf.text = title_text
    p = tf.paragraphs[0]
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Content blocks
    y_pos = 1.5
    for block in content_blocks:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos), Inches(9), Inches(block.get('height', 0.8)))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP

        p = tf.paragraphs[0]
        if isinstance(block['text'], str):
            p.text = block['text']
            p.font.size = Pt(block.get('size', 18))
            p.font.color.rgb = RGBColor(255, 255, 255)
            if block.get('bold'):
                p.font.bold = True
        else:
            for run_data in block['text']:
                run = p.add_run()
                run.text = run_data['text']
                run.font.size = Pt(run_data.get('size', 18))
                run.font.color.rgb = RGBColor(255, 255, 255)
                if run_data.get('bold'):
                    run.font.bold = True

        y_pos += block.get('height', 0.8) + 0.2

    # Footer
    footer = slide.shapes.add_textbox(Inches(0.5), Inches(5.1), Inches(9), Inches(0.3))
    tf = footer.text_frame
    tf.text = content_blocks[-1].get('footer', '')
    p = tf.paragraphs[0]
    p.font.size = Pt(10)
    p.font.italic = True
    p.font.color.rgb = LIGHT_GRAY

    return slide

def add_content_slide(prs, title_text, blocks):
    """Add a content slide with white background."""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # White background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    # Title
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.6))
    tf = title.text_frame
    tf.text = title_text
    p = tf.paragraphs[0]
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = DARK

    # Content blocks and footer
    footer_text = ''
    for block in blocks:
        if block.get('type') == 'text':
            box = slide.shapes.add_textbox(
                Inches(block['x']), Inches(block['y']),
                Inches(block['w']), Inches(block['h'])
            )
            tf = box.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.TOP

            p = tf.paragraphs[0]
            p.text = block['text']
            p.font.size = Pt(block.get('size', 14))
            p.font.color.rgb = block.get('color', DARK)
            if block.get('bold'):
                p.font.bold = True

        elif block.get('type') == 'box':
            shape = slide.shapes.add_shape(
                1,  # Rectangle
                Inches(block['x']), Inches(block['y']),
                Inches(block['w']), Inches(block['h'])
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = block.get('fill', LIGHT_GRAY)
            shape.line.color.rgb = block.get('line', GRAY)
            shape.line.width = Pt(1)

        if 'footer' in block:
            footer_text = block['footer']

    # Footer
    if footer_text:
        footer = slide.shapes.add_textbox(Inches(0.5), Inches(5.1), Inches(9), Inches(0.3))
        tf = footer.text_frame
        p = tf.paragraphs[0]
        p.text = footer_text
        p.font.size = Pt(10)
        p.font.italic = True
        p.font.color.rgb = GRAY

    return slide

# === SLIDE 1: User & Job ===
add_title_slide(prs, "User & Job", [
    {
        'text': [
            {'text': 'Job executor: ', 'bold': True},
            {'text': 'Học viên vừa xem xong slide bài học'}
        ],
        'height': 0.5
    },
    {
        'text': [
            {'text': 'Core JTBD: ', 'bold': True},
            {'text': 'Kiểm tra nhanh mức độ nhớ kiến thức trước khi hỏi tutor hoặc tiếp tục bài mới'}
        ],
        'height': 0.8
    },
    {
        'text': 'Pain Evidence',
        'size': 24,
        'bold': True,
        'height': 0.4
    },
    {
        'text': '''[TODO: Mining result]
VD: 67/200 chatlog (33.5%) học viên hỏi tutor "em có hiểu đúng không" hoặc "kiểm tra hiểu của em" ngay sau xem slide.

[TODO: Survey result]
VD: 18/22 học viên khảo sát (82%) xác nhận "muốn tự kiểm tra trước khi hỏi để không hỏi sai trọng tâm".''',
        'size': 14,
        'height': 1.5,
        'footer': 'CP1 Canvas → spec.md §1-2 + evidence log'
    }
])

# === SLIDE 2: Why This Feature ===
slide2 = add_content_slide(prs, "Vì sao chọn tính năng này", [
    {'type': 'text', 'x': 0.5, 'y': 1.3, 'w': 9, 'h': 0.4, 'text': 'Impact Comparison (3 ứng viên)', 'size': 20, 'bold': True, 'color': PRIMARY},
    {'type': 'text', 'x': 0.5, 'y': 1.9, 'w': 9, 'h': 2.8, 'text': '''[TODO: Bảng impact 3 ứng viên]

Cột: Ứng viên | Người gặp | Tần suất | Mỗi lần tốn | Build nổi? | Chọn?

1. Quiz từ slide: ~200 HV, mỗi buổi, 5-10' tìm câu hỏi, ✓, CHỌN
2. Tóm tắt buổi: [TODO], mỗi buổi, 15' viết notes, ✓, ✗ Evidence yếu
3. Chatbot logistics: [TODO], mỗi tuần, chờ TA, ✓, ✗ Ít người gặp''', 'size': 12},
    {'footer': 'spec.md §2 Impact Table + ứng viên đã loại'}
])

# === SLIDE 3: Solution & Demo ===
slide3 = add_content_slide(prs, "Giải pháp & Demo Live", [
    {'type': 'text', 'x': 0.5, 'y': 1.3, 'w': 4, 'h': 0.4, 'text': 'Lát cắt (1 câu)', 'size': 16, 'bold': True, 'color': PRIMARY},
    {'type': 'text', 'x': 0.5, 'y': 1.8, 'w': 4, 'h': 1.2, 'text': 'Học viên vừa xem xong slide → tải slide lên → AI tạo quiz 4 lựa chọn → làm quiz trong VLearn → xem đáp án + nguồn trích dẫn', 'size': 14},
    {'type': 'text', 'x': 0.5, 'y': 3.1, 'w': 4, 'h': 0.4, 'text': 'Automation: Conditional', 'size': 16, 'bold': True, 'color': PRIMARY},
    {'type': 'text', 'x': 0.5, 'y': 3.6, 'w': 4, 'h': 1.5, 'text': 'Chỉ tạo khi slide có đủ text rõ ràng. Slide toàn ảnh → từ chối + gợi ý.\n\nCost-of-error: Câu hỏi sai → học viên học sai kiến thức. Grounded refusal > confident hallucination.', 'size': 14},
    {'type': 'box', 'x': 5.2, 'y': 1.3, 'w': 4.3, 'h': 3.8},
    {'type': 'text', 'x': 5.2, 'y': 2.5, 'w': 4.3, 'h': 1, 'text': 'DEMO LIVE\n\n[Chạy prototype]\n\n1 case happy + 1 case chỗ khó', 'size': 14, 'color': GRAY},
    {'footer': 'spec.md §4 + LIVE không video'}
])

# === SLIDE 4: Results ===
slide4 = add_content_slide(prs, "Kết quả đo", [
    {'type': 'text', 'x': 0.5, 'y': 1.3, 'w': 4.5, 'h': 0.4, 'text': 'Quality Bar (chốt 23:59 N1)', 'size': 18, 'bold': True, 'color': PRIMARY},
    {'type': 'text', 'x': 0.5, 'y': 1.8, 'w': 4.5, 'h': 1, 'text': '[TODO: Bar]\n\nVD: Đạt khi ≥75% qua golden set VÀ 0% câu hỏi bịa nguồn', 'size': 14},
    {'type': 'text', 'x': 0.5, 'y': 3, 'w': 4.5, 'h': 0.4, 'text': 'Kết quả thực tế', 'size': 18, 'bold': True, 'color': PRIMARY},
    {'type': 'text', 'x': 0.5, 'y': 3.5, 'w': 4.5, 'h': 1, 'text': '[TODO: %]\n\nVD: 18/24 (75%) - ĐẠT\nVD: 16/24 (67%) - CHƯA ĐẠT', 'size': 14},
    {'type': 'box', 'x': 5.5, 'y': 1.3, 'w': 4, 'h': 3.5},
    {'type': 'text', 'x': 5.7, 'y': 1.5, 'w': 3.6, 'h': 0.4, 'text': 'Failure đáng kể nhất', 'size': 16, 'bold': True, 'color': PRIMARY},
    {'type': 'text', 'x': 5.7, 'y': 2, 'w': 3.6, 'h': 2.5, 'text': '[TODO: Case fail + phân tích]\n\nVD: 3/4 slide toàn ảnh vẫn gen → bịa.\nNguyên nhân: Threshold 50→200 chars', 'size': 13},
    {'footer': 'spec.md §7 + eval/runs/ — GHI NHẬN TRUNG THỰC'}
])

# === SLIDE 5: User Feedback ===
slide5 = add_content_slide(prs, "User thật nói gì", [
    {'type': 'box', 'x': 0.5, 'y': 1.3, 'w': 4.3, 'h': 1.5, 'line': PRIMARY},
    {'type': 'text', 'x': 0.7, 'y': 1.5, 'w': 3.9, 'h': 1.1, 'text': '[TODO: Quote 1]\n\nVD: "Slide toàn diagram, nó báo không đủ text — đúng." — Minh, HV B3', 'size': 13},
    {'type': 'box', 'x': 5.2, 'y': 1.3, 'w': 4.3, 'h': 1.5, 'line': PRIMARY},
    {'type': 'text', 'x': 5.4, 'y': 1.5, 'w': 3.9, 'h': 1.1, 'text': '[TODO: Quote 2]\n\nVD: "Trích dẫn trang giúp mình quay lại xem ngay." — Hương', 'size': 13},
    {'type': 'text', 'x': 0.5, 'y': 3.1, 'w': 9, 'h': 0.4, 'text': 'Thay đổi từ feedback', 'size': 20, 'bold': True, 'color': PRIMARY},
    {'type': 'text', 'x': 0.5, 'y': 3.6, 'w': 9, 'h': 1.5, 'text': '[TODO: ≥1 thay đổi]\n\nVD: ĐÃ SỬA: Thêm nút Hủy rõ ràng.\nVD: GIỮ: One-at-a-time vì preserves momentum.', 'size': 13},
    {'footer': 'validation/ ≥5 người + Changelog'}
])

# === SLIDE 6: Next Week ===
add_title_slide(prs, "Nếu có thêm 1 tuần", [
    {
        'text': 'Ưu tiên từ feedback/failure',
        'size': 20,
        'bold': True,
        'height': 0.4
    },
    {
        'text': '''1. [TODO: Priority từ failure]
   VD: Nâng threshold 50→200 chars chặn slide image-heavy

2. [TODO: Priority từ validation]
   VD: Link "Xem nguồn" mở PDF tại đúng trang

3. [TODO: Backlog]
   VD: OCR tiếng Việt có dấu''',
        'size': 15,
        'height': 1.6
    },
    {
        'text': 'Bài học lớn nhất',
        'size': 20,
        'bold': True,
        'height': 0.4
    },
    {
        'text': '[TODO: Bài học từ case fail]\n\nVD: Khi AI refuse, PHẢI có actionable recovery — không gợi ý "làm gì tiếp" thì user bỏ luôn.',
        'size': 14,
        'height': 0.8,
        'footer': ''
    }
])

# Save
prs.save('demo-slides.pptx')
print("✓ Created demo-slides.pptx")
